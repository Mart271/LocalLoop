#!/usr/bin/env python3
"""EV-04: independently re-read Rust round trips, and test an exclusive Windows file lock.

Uses synthetic fixtures and temporary output files only. Never imported by product code.
"""
from __future__ import annotations

import argparse
import ctypes
import json
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from common import ROOT, digest, metadata, write_report


def xml(element) -> str | None:
    return ET.tostring(element, encoding="unicode") if element is not None else None


def snapshot(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheets = {}
    for sheet in workbook:
        cells = {}
        for row in sheet:
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                cells[cell.coordinate] = {
                    "value": str(cell.value), "type": cell.data_type,
                    "number_format": cell.number_format,
                    "font": xml(cell.font.to_tree()), "fill": xml(cell.fill.to_tree()),
                    "border": xml(cell.border.to_tree()), "alignment": xml(cell.alignment.to_tree()),
                    "comment": (cell.comment.text, cell.comment.author) if cell.comment else None,
                }
        sheets[sheet.title] = {
            "state": sheet.sheet_state, "freeze_panes": sheet.freeze_panes,
            "cells": cells,
            "column_widths": {k: v.width for k, v in sheet.column_dimensions.items()},
            "tables": {k: xml(sheet.tables[k].to_tree()) for k in sheet.tables},
            "validation": xml(sheet.data_validations.to_tree()),
            "conditional_formatting": {
                str(k.sqref): [xml(r.to_tree()) for r in rules]
                for k, rules in sheet.conditional_formatting._cf_rules.items()
            },
            "chart_count": len(sheet._charts),
            "chart_types": [type(c).__name__ for c in sheet._charts],
        }
    result = {"sheets": sheets, "defined_names": {k: v.attr_text for k, v in workbook.defined_names.items()}}
    workbook.close()
    return result


def differences(before, after, prefix="") -> list[str]:
    if isinstance(before, dict) and isinstance(after, dict):
        result = []
        for key in sorted(before.keys() | after.keys()):
            path = f"{prefix}/{key}"
            if key not in before or key not in after:
                result.append(path)
            else:
                result.extend(differences(before[key], after[key], path))
        return result
    return [] if before == after else [prefix]


def lock_test(path: Path) -> dict:
    from ctypes import wintypes
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p,
                                  wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    kernel.CreateFileW.restype = wintypes.HANDLE
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    handle = kernel.CreateFileW(str(path), 0x80000000, 0, None, 3, 0x80, None)
    if handle == ctypes.c_void_p(-1).value:
        raise ctypes.WinError(ctypes.get_last_error())
    try:
        try:
            with path.open("r+b"):
                blocked = False
        except PermissionError:
            blocked = True
    finally:
        kernel.CloseHandle(handle)
    with path.open("r+b"):
        reopened = True
    return {"exclusive_lock_blocks_write": blocked, "reopens_after_release": reopened,
            "method": "CreateFileW share mode 0; simulates an exclusive holder, not an Excel UI test"}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--binary", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    report = {"environment": metadata(), "library": "umya-spreadsheet 3.1.0 (MIT)", "cases": [],
              "limitations": ["No Excel rendering or recalculation test", "Chart count/type checked; chart appearance unverified",
                              "Finite fixtures do not establish support for arbitrary XLSX packages"]}
    folder = ROOT / "tests/fixtures/generated/F6-spreadsheets"
    with tempfile.TemporaryDirectory(prefix="localloop-xlsx-") as temporary:
        output_dir = Path(temporary)
        for name in ("Invoices.xlsx", "Invoices-formatted.xlsx"):
            source = folder / name
            original = digest(source)
            before = snapshot(source)
            for operation in ("roundtrip", "edit", "preserving-edit"):
                output = output_dir / f"{operation}-{name}"
                completed = subprocess.run([str(args.binary.resolve()), str(source), str(output), operation],
                                           capture_output=True, text=True, timeout=60)
                case = {"file": name, "operation": operation, "exit_code": completed.returncode,
                        "source_unchanged": digest(source) == original, "measure": completed.stdout.strip()}
                if completed.returncode == 0:
                    after = snapshot(output)
                    if operation != "roundtrip":
                        case["edit_verified"] = after["sheets"]["Invoices"]["cells"]["H2"]["value"] == "spike-updated.pdf"
                        after["sheets"]["Invoices"]["cells"]["H2"]["value"] = before["sheets"]["Invoices"]["cells"]["H2"]["value"]
                    case["unexpected_semantic_changes"] = differences(before, after)
                    with zipfile.ZipFile(source) as old, zipfile.ZipFile(output) as new:
                        case["removed_package_parts"] = sorted(set(old.namelist()) - set(new.namelist()))
                        case["changed_package_parts"] = sorted(n for n in set(old.namelist()) & set(new.namelist()) if old.read(n) != new.read(n))
                else:
                    case["error"] = completed.stderr[-2000:]
                report["cases"].append(case)
        report["lock"] = lock_test(folder / "Invoices.xlsx")
        # Record unsafe features the future adapter must detect before a parser/rewriter runs.
        report["macro_fixture"] = {"path": "Invoices-macros.xlsm", "present": (folder / "Invoices-macros.xlsm").exists()}
    write_report(args.output, report)
    print(json.dumps(report, indent=2))
    if not all(c["source_unchanged"] and c["exit_code"] == 0 for c in report["cases"]):
        raise SystemExit(1)
    for case in report["cases"]:
        if case["operation"] == "preserving-edit" and (case["unexpected_semantic_changes"] or case["removed_package_parts"] or case["changed_package_parts"] != ["xl/worksheets/sheet1.xml"] or not case["edit_verified"]):
            raise SystemExit("package-preserving edit failed independent verification")
    if not all(v for k, v in report["lock"].items() if k != "method"):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
