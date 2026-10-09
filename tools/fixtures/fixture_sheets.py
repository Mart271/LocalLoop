"""Synthetic spreadsheets for set F6 (LL-012, EV-04).

openpyxl stamps the current time into document properties and zip entries, so every workbook is
rewritten with fixed timestamps to keep the output reproducible.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

FIXED_TIME = datetime(2026, 1, 1, 0, 0, 0)
ZIP_TIME = (1980, 1, 1, 0, 0, 0)

COLUMNS = ["Vendor", "Invoice No", "Invoice Date", "Currency", "Subtotal", "VAT", "Total", "Source File"]

# Rows already logged before any LocalLoop run (fictitious, dated 2025 so they never collide with F1).
HISTORIC_ROWS: list[list[object]] = [
    ["Bluefin Office Supplies Co.", "BOS-2025-0187", date(2025, 11, 3), "PHP", Decimal("4210.00"), Decimal("505.20"),
     Decimal("4715.20"), "BOS-2025-0187.pdf"],
    ["Maraval Printing Services", "MPS-2025-0412", date(2025, 12, 9), "USD", Decimal("880.50"), Decimal("70.44"),
     Decimal("950.94"), "MPS-2025-0412.pdf"],
    ["Tanglewood Logistics Ltd.", "TGL-2025-0093", date(2025, 12, 18), "EUR", Decimal("1320.00"), Decimal("264.00"),
     Decimal("1584.00"), "TGL-2025-0093.pdf"],
]


def _stamp(wb: Workbook) -> None:
    wb.properties.creator = "LocalLoop fixture generator"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.properties.lastModifiedBy = "LocalLoop fixture generator"


def normalize_zip(path: Path, extra: dict[str, bytes] | None = None, replace: dict[str, bytes] | None = None) -> None:
    """Rewrite a zip with fixed timestamps and stored entries (and optional added or replaced parts)."""
    with zipfile.ZipFile(path) as source:
        parts = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    replace = replace or {}
    parts = [(name, replace.get(name, data)) for name, data in parts]
    # openpyxl writes the save time into dcterms:modified regardless of the workbook properties.
    parts = [
        (name, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-01-01T00:00:00Z\g<2>", data))
        if name == "docProps/core.xml" else (name, data)
        for name, data in parts
    ]
    parts += sorted((extra or {}).items())
    # Stored (uncompressed) entries: deflate output differs between zlib and zlib-ng.
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_STORED) as target:
        for name, data in parts:
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_STORED
            info.external_attr = 0o644 << 16
            info.create_system = 3  # record a fixed creator OS (Python uses 0 on Windows, 3 elsewhere)
            target.writestr(info, data)


def _save(wb: Workbook, path: Path) -> None:
    _stamp(wb)
    wb.save(path)
    normalize_zip(path)


def plain_workbook(path: Path, columns: list[str] = COLUMNS, sheet: str = "Invoices") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(columns)
    for row in HISTORIC_ROWS:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row][: len(columns)])
    _save(wb, path)


def formatted_workbook(path: Path) -> None:
    """A workbook with the features EV-04 must check for round-trip fidelity."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(COLUMNS)
    for row in HISTORIC_ROWS:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[2].number_format = "yyyy-mm-dd"
        for cell in row[4:7]:
            cell.number_format = "#,##0.00"
    for letter, width in zip("ABCDEFGH", (32, 16, 14, 10, 14, 12, 14, 24)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    table = Table(displayName="InvoiceTable", ref=f"A1:H{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.conditional_formatting.add(
        f"G2:G{ws.max_row}", CellIsRule(operator="greaterThan", formula=["10000"], fill=PatternFill("solid", fgColor="FFC7CE"))
    )
    validation = DataValidation(type="list", formula1='"PHP,USD,EUR"', allow_blank=False)
    ws.add_data_validation(validation)
    validation.add("D2:D1000")
    ws["B2"].comment = Comment("Logged manually before LocalLoop (synthetic).", "fixture")

    summary = wb.create_sheet("Summary")
    summary["A1"] = "Invoices logged"
    summary["B1"] = "=COUNTA(Invoices!B:B)-1"
    summary["A2"] = "Total of all invoices"
    summary["B2"] = "=SUM(Invoices!G:G)"
    summary["B2"].number_format = "#,##0.00"
    summary["A3"] = "Default VAT rate"
    summary["B3"] = 0.12
    wb.defined_names["DefaultVatRate"] = DefinedName("DefaultVatRate", attr_text="Summary!$B$3")
    chart = BarChart()
    chart.title = "Invoice totals"
    chart.add_data(Reference(ws, min_col=7, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=2, max_row=ws.max_row))
    summary.add_chart(chart, "D2")

    lists = wb.create_sheet("Lists")
    for value in ("PHP", "USD", "EUR"):
        lists.append([value])
    lists.sheet_state = "hidden"
    _save(wb, path)


VBA_CONTENT_TYPE = b"application/vnd.ms-excel.sheet.macroEnabled.main+xml"
XLSX_CONTENT_TYPE = b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"


def macro_workbook(path: Path) -> None:
    """A macro-enabled package (.xlsm) that LocalLoop must refuse to modify (A-05).

    openpyxl cannot author VBA, so the package declares the macro-enabled content type and carries a
    placeholder `xl/vbaProject.bin` part. It contains no executable code.
    """
    plain_workbook(path)
    with zipfile.ZipFile(path) as source:
        content_types = source.read("[Content_Types].xml")
        workbook_rels = source.read("xl/_rels/workbook.xml.rels")
    content_types = content_types.replace(XLSX_CONTENT_TYPE, VBA_CONTENT_TYPE).replace(
        b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>'
    )
    workbook_rels = workbook_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
        b'Target="vbaProject.bin"/></Relationships>',
    )
    normalize_zip(
        path,
        extra={"xl/vbaProject.bin": b"LOCALLOOP-SYNTHETIC-PLACEHOLDER: not a real VBA project\n"},
        replace={"[Content_Types].xml": content_types, "xl/_rels/workbook.xml.rels": workbook_rels},
    )


def csv_ledger(path: Path, bom: bool = False) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(COLUMNS)
    for row in HISTORIC_ROWS:
        writer.writerow([v.isoformat() if isinstance(v, date) else str(v) for v in row])
    data = buffer.getvalue().encode("utf-8")
    path.write_bytes((b"\xef\xbb\xbf" if bom else b"") + data)
